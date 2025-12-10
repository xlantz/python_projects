import pandas as pd
from googletrans import Translator
from pypinyin import pinyin, Style
from tqdm import tqdm
import re
import os
from langdetect import detect
import spacy
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

# ----------------------------
# Load CC-CEDICT dictionary
# ----------------------------
def load_cedict(file_path="cedict_ts.u8"):
    cedict = {}
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                if line.startswith("#"):
                    continue
                match = re.match(r"(\S+) (\S+) \[(.+?)\] /(.+)/", line)
                if match:
                    traditional, simplified, py, definitions = match.groups()
                    first_def = re.split(r";|\(", definitions)[0].strip()
                    cedict[simplified] = {"pinyin": py, "english": first_def}
    except FileNotFoundError:
        print("⚠️ CC-CEDICT file not found, Chinese translations will use Google Translate only.")
    return cedict

cedict = load_cedict()

# ----------------------------
# Load spaCy models for supported languages
# ----------------------------
spacy_models = {}
lang_model_map = {
    "de": "de_core_news_sm",
    "es": "es_core_news_sm",
    "fr": "fr_core_news_sm",
    "it": "it_core_news_sm",
    "pl": "pl_core_news_sm",
    "ru": "ru_core_news_sm"
}
for code, model in lang_model_map.items():
    try:
        spacy_models[code] = spacy.load(model)
    except Exception:
        print(f"⚠️ spaCy model for {code} not found. Lemmas for this language will default to original words.")

# ----------------------------
# Helper functions
# ----------------------------
def contains_chinese(text):
    return any('\u4e00' <= ch <= '\u9fff' for ch in str(text))

def get_cedict_data(word):
    entry = cedict.get(word)
    if entry:
        return entry["pinyin"], entry["english"], word
    return None, None, word

def get_lemma(word, lang_code):
    if lang_code in spacy_models:
        doc = spacy_models[lang_code](word)
        return " ".join([token.lemma_ for token in doc])
    return word

# ----------------------------
# Setup
# ----------------------------
print("🌍 Multilingual Enricher — Translate & Enrich Word Lists")
input_path = input("Enter the path of your Excel file: ").strip('"')
output_name = input("Enter a name for the output file (extension optional): ").strip('"')

# Default extension
if not any(output_name.endswith(ext) for ext in [".xlsx", ".xls", ".csv", ".txt"]):
    output_name += ".xlsx"

# Save in same folder as input
input_dir = os.path.dirname(input_path)
output_path = os.path.join(input_dir, output_name)

num_words = int(input("How many words below the header do you want to process? (e.g., 200): "))

print("\nDo you want English translations from:")
print("[1] CC-CEDICT (accurate dictionary style, only for Chinese)")
print("[2] Google Translate (all languages, simpler)")
choice = input("Enter 1 or 2: ").strip()

# Load input file
df = pd.read_excel(input_path)
column_name = df.columns[0]
words = df[column_name].dropna().head(num_words)

translator = Translator()

# ----------------------------
# Output containers
# ----------------------------
translations = []
pinyins = []
lemmas = []
errors = []

# ----------------------------
# Main loop with progress bar
# ----------------------------
print("\n⚙️ Starting... Please wait.\n")
for word in tqdm(words, desc="Processing words", ncols=100):
    try:
        lang_code = detect(str(word))
        lemma = word
        py = ""

        if contains_chinese(word):
            # Mandarin
            cedict_py, cedict_trans, cedict_lemma = get_cedict_data(word)
            if choice == "1" and cedict_trans:
                translation = cedict_trans
            else:
                translation = translator.translate(word, src='zh-cn', dest='en').text

            # Pinyin
            if cedict_py:
                py = cedict_py
            else:
                py_list = pinyin(word, style=Style.TONE)
                py = " ".join([s[0] for s in py_list])

            lemma = cedict_lemma

        else:
            # Non-Chinese
            translation = translator.translate(word, dest='en').text
            py = ""
            lemma = get_lemma(word, lang_code)

        translations.append(translation)
        pinyins.append(py)
        lemmas.append(lemma)

    except Exception as e:
        translations.append("")
        pinyins.append("")
        lemmas.append("")
        errors.append({"Word": word, "Error": str(e)})

# ----------------------------
# Create output DataFrame
# ----------------------------
output_df = pd.DataFrame({
    column_name: words,
    "English_Translation": translations,
    "Pinyin": pinyins,
    "Lemma": lemmas
})

# ----------------------------
# Save to Excel with highlighting for untranslated words
# ----------------------------
with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    output_df.to_excel(writer, index=False, sheet_name="Enriched")
    if errors:
        pd.DataFrame(errors).to_excel(writer, index=False, sheet_name="Error Translations")

# Highlight cells where English translation equals original word
wb = load_workbook(output_path)
ws = wb["Enriched"]
yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

header = [cell.value for cell in ws[1]]
word_col = header.index(column_name) + 1
eng_col = header.index("English_Translation") + 1

for row in range(2, ws.max_row + 1):
    word_cell = ws.cell(row=row, column=word_col)
    eng_cell = ws.cell(row=row, column=eng_col)
    if str(word_cell.value).strip().lower() == str(eng_cell.value).strip().lower():
        word_cell.fill = yellow_fill
        eng_cell.fill = yellow_fill

wb.save(output_path)
print(f"\n✅ Translation complete. Saved to: {output_path} (untranslated words highlighted in yellow)")
